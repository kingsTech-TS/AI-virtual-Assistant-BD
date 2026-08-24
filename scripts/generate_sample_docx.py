import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

os.makedirs('sample_documents', exist_ok=True)

doc = docx.Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = p_title.add_run('APEX FEDERAL UNIVERSITY OF TECHNOLOGY\nUNDERGRADUATE ADMISSIONS HANDBOOK')
run_title.font.name = 'Arial'
run_title.font.size = Pt(16)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = p_sub.add_run('Official Admissions Guidelines & Policy (2026/2027 Academic Session)\nDocument Ref: AFUTS/ADM/HB/2026-V1')
run_sub.font.name = 'Arial'
run_sub.font.size = Pt(10)
run_sub.font.italic = True
run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

doc.add_heading('1. Overview & Admission Philosophy', level=1)
doc.add_paragraph(
    'The Admissions Board of Apex Federal University operates an open, merit-driven, and transparent admissions '
    'framework designed to enroll qualified undergraduate candidates into degree programs across Science, '
    'Engineering, Management, and Technology disciplines.'
)

doc.add_heading('2. General Undergraduate Admission Requirements (100-Level)', level=1)
doc.add_paragraph('Candidates seeking admission into 100-Level degree programs must satisfy the following criteria:')
doc.add_paragraph('• Minimum Age: Candidates must be at least 16 years of age by October 31 of the entry year.', style='List Bullet')
doc.add_paragraph('• O-Level Qualifications: Minimum of five (5) credit passes in relevant SSCE (WAEC/NECO/NABTEB) subjects in not more than two (2) sittings.', style='List Bullet')
doc.add_paragraph('• Compulsory Core: English Language and Mathematics are mandatory for all undergraduate courses.', style='List Bullet')
doc.add_paragraph('• UTME Score: Minimum score of 200 in the Unified Tertiary Matriculation Examination (UTME) for Sciences and Engineering; 180 for Social Sciences.', style='List Bullet')
doc.add_paragraph('• Institution Choice: The candidate must choose the institution as their First Choice on the JAMB portal.', style='List Bullet')

doc.add_heading('3. Departmental Requirements & Subject Combinations', level=1)
table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
headers = ['Department', 'O-Level Requirements', 'UTME Combination', 'Cut-Off Score']
for i, head in enumerate(headers):
    hdr_cells[i].text = head
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True

data = [
    ('Computer Science (CSC)', 'English, Math, Physics, Chemistry, Biology/F.Math', 'English, Math, Physics, Chemistry', '220 (UTME) + 65% CBT'),
    ('Physics (PHY)', 'English, Math, Physics, Chemistry, Biology', 'English, Physics, Math, Chemistry', '200 (UTME) + 55% CBT'),
    ('Mathematics (MAT)', 'English, Math, Physics, Further Math/Economics', 'English, Math, Physics, Chemistry', '200 (UTME) + 50% CBT'),
    ('Mechanical Engineering', 'English, Math, Physics, Chemistry, Technical Drawing', 'English, Math, Physics, Chemistry', '230 (UTME) + 70% CBT'),
]

for row in data:
    r_cells = table.add_row().cells
    for j, val in enumerate(row):
        r_cells[j].text = val

doc.add_heading('4. Direct Entry (200-Level) Criteria', level=1)
doc.add_paragraph('• IJMB / JUPEB / Cambridge GCE: Minimum of 10 points across three relevant Advanced Level subjects.', style='List Bullet')
doc.add_paragraph('• National Diploma (ND): Upper Credit (minimum CGPA 3.00/4.00) from an accredited institution in a related discipline.', style='List Bullet')
doc.add_paragraph('• Higher National Diploma (HND): Upper Credit holders may be considered for 200-Level placement.', style='List Bullet')

doc.add_heading('5. Application, Screening & Acceptance Procedure', level=1)
doc.add_paragraph(
    '1. Register online at admissions.afuts.edu.ng and pay the screening fee of N2,000.\n'
    '2. Upload verified O-Level results, UTME slips, and statutory identification documents.\n'
    '3. Attend the Computer-Based Screening Test at the University ICT Center on your scheduled date.\n'
    '4. Accept the admission offer on JAMB CAPS and pay the N35,000 acceptance fee within two weeks.\n'
    '5. Complete physical departmental clearance with original documents to receive your matriculation number.'
)

doc.add_heading('6. Helpdesk & Inquiries', level=1)
doc.add_paragraph(
    'Admissions Helpdesk: admissions-support@afuts.edu.ng | +234 (0) 800-AFUTS-ADM\n'
    'In-Portal Academic AI Chatbot available 24/7 for automated inquiries and support ticket creation.'
)

doc.save('sample_documents/undergraduate_admissions_guide_2026.docx')
print('DOCX generated successfully.')
